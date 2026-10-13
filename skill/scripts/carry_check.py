#!/usr/bin/env python3
"""再評価で、前回の台帳の ID が今回の台帳に引き継がれているかを確かめる。

    python carry_check.py <前回の台帳.xlsx> <今回の台帳.xlsx>

再評価では、前回の指摘（S-x・N-x）と未確認事項（U-x）を 1 件ずつ突き合わせ、行を消さずに残す
（references/11-reassessment.md の 2・3 節）。手で写すと、行がどこにも現れないまま消えたり、
同じ番号が別の意味に使われたりする。実際に、前回の指摘 1 件と未確認事項 8 件が理由の記載なく消え、
未確認事項の番号が振り直されて、前回の番号を参照した記述が別の項目を指した回があった。

見るもの:
  1. 前回の台帳で行のある ID のうち、今回の台帳に行の無いもの（消えた行）
  2. 今回の台帳の文中で参照しているのに、どのシートにも行の無い ID（参照切れ）
  3. 前回と今回で同じ ID なのに、行の中身が大きく違うもの（番号の使い回しの疑い。見比べて確かめる）
  4. 前回の未確認事項（U-x）のうち、今回の行の「行き先」が空か、確定・未確認のまま・取り下げの
     どれでもないもの。今回の台帳にその列が無ければ（古い版の雛形のまま）、そのことを出す

「行がある」は、どれかのシートの A 列にその ID だけが入っていること。参照は、文中に現れる
「英大文字 1〜2 字-数字」の形（S-08・U-4・S-22a）。「02 の A-1」のような資料の節の参照と、
CVE・CWE の番号は数えない。

シートは名前で探さない。前回の台帳は古い版の雛形で作られていることが多く、未確認事項のシートの名前も
版によって違う（「5_未確認事項」「7_残る未確認事項」「7_未確認事項」）。行き先の列は、見出しの文字で探す。

終了コード: 0 = 1・2・4 が無い / 2 = 1・2・4 のどれかがある / 1 = 使い方の誤り・読めない。
3 は見比べる候補として出すだけで、終了コードには数えない（文言を直しただけの行も出るため）。
"""
import os
import re
import sys

ID_CELL = re.compile(r"^([A-Z]{1,2}-[0-9]{1,3}[a-z]?)$")
# 前に英数字・ハイフンが付かない（CVE-2025-… や GHSA-… の一部を拾わない）
ID_REF = re.compile(r"(?<![A-Za-z0-9-])([A-Z]{1,2}-[0-9]{1,3}[a-z]?)(?![0-9A-Za-z-])")
# 「02 の A-1」「01 の B-3」のような、スキルの資料の節の参照
SECTION_REF = re.compile(r"[0-9]{2}[ 　]*の[ 　]*$")
# 未確認事項の行き先の列の見出しと値（make_register.py の UNKNOWN_COLS・DESTS と同じ）
DEST_HEAD = "行き先"
DESTS = ("確定", "未確認のまま", "取り下げ")
NO_COLUMN = object()   # 行のあるシートに行き先の列が無い


def load(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("openpyxl が要る（make_register.py と同じ。python3 -m pip install openpyxl）", file=sys.stderr)
        raise SystemExit(1)
    if not os.path.isfile(path):
        print(f"開けない: {path}", file=sys.stderr)
        raise SystemExit(1)
    try:
        return load_workbook(path, data_only=False)
    except Exception as e:  # 壊れた xlsx・別の形式
        print(f"台帳として読めない: {path}（{e}）", file=sys.stderr)
        raise SystemExit(1)


def rows_and_refs(wb):
    """ID の行（ID → (シート名, 行の文)）と、文中の参照（ID → [シート名!セル]）と、
    未確認事項の行き先（U-x → 値。列が無ければ NO_COLUMN）を返す"""
    rows, refs, dests = {}, {}, {}
    for ws in wb.worksheets:
        dest_col = None   # 見出しの行より下で使う
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() == DEST_HEAD:
                    dest_col = c.column - 1
            first = row[0].value if row else None
            if isinstance(first, str) and ID_CELL.match(first.strip()):
                rid = first.strip()
                text = " ".join(str(c.value) for c in row[1:] if isinstance(c.value, str))
                if rid not in rows and rid.startswith("U-"):
                    v = row[dest_col].value if dest_col is not None and dest_col < len(row) else None
                    dests[rid] = NO_COLUMN if dest_col is None else ("" if v is None else str(v).strip())
                rows.setdefault(rid, (ws.title, text))
            for c in row:
                v = c.value
                if not isinstance(v, str) or (c is row[0] and ID_CELL.match(v.strip())):
                    continue
                if v.startswith("="):  # 数式の中の文字列（集計の条件）は参照ではない
                    continue
                for m in ID_REF.finditer(v):
                    if SECTION_REF.search(v[:m.start()]):
                        continue
                    refs.setdefault(m.group(1), []).append(f"{ws.title}!{c.coordinate}")
    return rows, refs, dests


def overlap(a, b):
    """2 つの行の文の重なり（文字の 2 字組の重なりを、少ない方の数で割る。0〜1）。
    【前回から未解消】のような印と空白は除く。書き直しただけの行は 0.27 以上、番号を別の意味に
    使い回した行は 0.16 以下だった（実案件の 1 回の再評価で測った値）。どちらかが空なら比べない"""
    def grams(s):
        s = re.sub(r"\s+", "", re.sub(r"【[^】]*】", "", s))[:400]
        return {s[i:i + 2] for i in range(len(s) - 1)}
    ga, gb = grams(a), grams(b)
    if not ga or not gb:
        return 1.0
    return len(ga & gb) / min(len(ga), len(gb))


def short(s, n=70):
    s = " ".join(s.split())
    return s if len(s) <= n else s[:n] + "…"


def main():
    if len(sys.argv) != 3:
        print(__doc__.split("\n\n")[1], file=sys.stderr)
        raise SystemExit(1)
    prev_rows, _, _ = rows_and_refs(load(sys.argv[1]))
    cur_rows, cur_refs, cur_dests = rows_and_refs(load(sys.argv[2]))

    missing = sorted(set(prev_rows) - set(cur_rows), key=sortkey)
    dangling = sorted(set(cur_refs) - set(cur_rows), key=sortkey)
    changed = [rid for rid in sorted(set(prev_rows) & set(cur_rows), key=sortkey)
               if overlap(prev_rows[rid][1], cur_rows[rid][1]) < 0.2]
    # 前回の U-x を今回も残したなら、行き先を選んでいるか（11 の 3 節）
    carried_u = [rid for rid in sorted(set(prev_rows) & set(cur_rows), key=sortkey) if rid.startswith("U-")]
    no_column = [rid for rid in carried_u if cur_dests.get(rid) is NO_COLUMN]
    no_dest = [rid for rid in carried_u if rid not in no_column and cur_dests.get(rid) not in DESTS]

    print(f"=== 前回の台帳の ID の引き継ぎ（前回 {len(prev_rows)} 件・今回 {len(cur_rows)} 件）===")
    print(f"--- 1. 前回にあって今回に無い ID（{len(missing)} 件。行を消さず、行き先を書いて残す。11 の 2・3 節）---")
    for rid in missing:
        print(f"  {rid}  前回の {prev_rows[rid][0]}: {short(prev_rows[rid][1])}")
    print(f"--- 2. 今回の台帳で参照しているのに行の無い ID（{len(dangling)} 件）---")
    for rid in dangling:
        where = cur_refs[rid]
        print(f"  {rid}  {', '.join(where[:3])}{' ほか' if len(where) > 3 else ''}")
    print(f"--- 3. 同じ ID で中身が大きく違うもの（{len(changed)} 件。番号を別の意味に使っていないか見比べる）---")
    for rid in changed:
        print(f"  {rid}  前回: {short(prev_rows[rid][1], 50)}")
        print(f"  {' ' * len(rid)}  今回: {short(cur_rows[rid][1], 50)}")
    print(f"--- 4. 前回の未確認事項で、今回の行き先（{'・'.join(DESTS)}）を選んでいないもの"
          f"（{len(no_dest) + len(no_column)} 件。11 の 3 節）---")
    for rid in no_dest:
        v = cur_dests.get(rid)
        print(f"  {rid}  今回の {cur_rows[rid][0]}: 行き先が" + (f"「{v}」" if v else "空"))
    if no_column:
        sheets = "・".join(sorted({cur_rows[rid][0] for rid in no_column}))
        print(f"  {' '.join(no_column)}  今回の {sheets} に「{DEST_HEAD}」の列が無い"
              "（古い版の雛形のまま。今の版の make_register.py で作り直して前回の行を移す。11 の 0 節）")
    if missing or dangling or no_dest or no_column:
        print("※ 1・2・4 が 0 件になるまで直す。3 は、文言を直しただけなら問題ない")
        raise SystemExit(2)
    print("※ 1・2・4 は 0 件。3 は、文言を直しただけなら問題ない")


def sortkey(rid):
    m = re.match(r"([A-Z]+)-([0-9]+)([a-z]?)", rid)
    return (m.group(1), int(m.group(2)), m.group(3)) if m else (rid, 0, "")


if __name__ == "__main__":
    main()
