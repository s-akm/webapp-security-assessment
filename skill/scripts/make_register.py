#!/usr/bin/env python3
"""指摘台帳（xlsx）の雛形を生成する。

    python make_register.py <出力先.xlsx> [--service "サービス名"] [--date 2026-01-15]
                            [--skill-version 2.18.0] [--frameworks none|owasp|full]
                            [--owasp 2025|2021] [--api] [--card] [--force]

出力先が既にあれば止まる（書きかけの台帳を雛形で潰さないため）。上書きするなら --force。

--skill-version は、評価に使ったスキルの版（配布物の VERSION の値）を総合評価の「評価の前提」に入れる。
このスクリプトは skill/ の外にある VERSION を読めないので、渡されなければ空欄にする。

指摘事項一覧の 1 行には、判定（問題あり／問題なし／判断保留）・優先度（P0〜P4、問題なしは —）・
状態（未対応／対応中／クローズ（解消）／クローズ（該当なし）／見送り）を別の列で持つ。
いずれも入力規則で一覧から選ぶ。対応が要る件数は「判定が問題ありで、状態がクローズでも
見送りでもないもの」を COUNTIFS で数える（references/04-findings-register.md）。

--frameworks で枠組みへの当てはめシートの構成を選ぶ。

    none   6 枚。個別の指摘だけを扱う。対外説明が要らない案件向け
    owasp  7 枚（既定）。OWASP Top 10 を 1 枚にまとめる
    full   9 枚。個人情報の安全管理措置・OWASP・IPA 非機能要求グレードを
           それぞれ独立させる。対外説明や水準の提示が要る案件向け

--api は API が主体の構成で使う。一般の Top 10 と重なるため両方は並べず、OWASP のシートを
同じ番号のまま API Security Top 10 のシートに置き換える（owasp なら 7_API_Top10、
full なら 4_API_Top10）。none では 7_API_Top10 が 1 枚増える。

--card はカード決済を扱うとき（決済代行の画面へ遷移する・埋め込む構成を含む）だけ付ける。
末尾に「<次の番号>_カード決済」が 1 枚増える（owasp なら 8_カード決済）。常設はしない。

件数と工数の合計は数式で入っているので、指摘の行を足せば自動で追従する。
参照範囲は 200 行まで取ってあるため、行を足すときに数式へ手を入れる必要はない。

各シートには例示行を 1 行だけ入れてある（薄い黄色の網掛け）。書き方の見本なので、
実際の指摘を書き始めるときに上書きするか削除する。

注意: openpyxl は数式を文字列として書き込むだけで、計算結果は保持しない。
表計算ソフトで開くまで合計欄は空に見える。これは正常な挙動。
"""

import argparse
import datetime as _dt
import os
import sys


def _externally_managed():
    """pip install がこの python に対して拒否されるか（PEP 668）。

    Homebrew の python3.13 / 3.14 や Debian 系の python3 は EXTERNALLY-MANAGED を置いており、
    仮想環境の外で pip install すると「externally-managed-environment」で失敗する。
    その環境に pip install を案内すると、利用者は案内どおりにして失敗する。
    """
    if sys.prefix != getattr(sys, "base_prefix", sys.prefix):
        return False  # 仮想環境の中なら pip で入る
    try:
        import sysconfig
        return os.path.exists(os.path.join(sysconfig.get_path("stdlib"), "EXTERNALLY-MANAGED"))
    except Exception:  # noqa: BLE001
        return False


try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # noqa: BLE001
    import shutil
    import subprocess

    _me = os.path.abspath(__file__)
    print("openpyxl が見つからない。次のいずれかで導入する。", file=sys.stderr)
    if _externally_managed():
        print("  この python は外部管理（PEP 668）のため、pip install は失敗する。", file=sys.stderr)
    else:
        print(f"  {sys.executable} -m pip install openpyxl", file=sys.stderr)
    # 仮想環境は評価対象のリポジトリの中に作らない（対象を書き換えてしまう）。スキルの外に置く
    print("  仮想環境を作る（評価対象のリポジトリの外に）:", file=sys.stderr)
    print('    python3 -m venv "$HOME/.cache/wsa-venv" && "$HOME/.cache/wsa-venv/bin/pip" install openpyxl', file=sys.stderr)
    print(f'    "$HOME/.cache/wsa-venv/bin/python" {_me} <出力先.xlsx>', file=sys.stderr)
    print("  uv があれば、入れずにその場で使える:", file=sys.stderr)
    print(f"    uv run --with openpyxl python {_me} <出力先.xlsx>", file=sys.stderr)
    print("  Debian / Ubuntu の python3 なら: sudo apt install python3-openpyxl", file=sys.stderr)
    # 別の python に入っていることが多い。探して案内する。
    for _cand in ("/usr/bin/python3", "python3.14", "python3.13", "python3.12",
                  "python3.11", "python3.10", "python3.9"):
        _path = _cand if _cand.startswith("/") else shutil.which(_cand)
        if not _path or not os.path.exists(_path) or os.path.realpath(_path) == os.path.realpath(sys.executable):
            continue
        _r = subprocess.run(
            [_path, "-c", "import openpyxl"], capture_output=True, check=False
        )
        if _r.returncode == 0:
            print(f"  すでに入っている python がある: {_path}", file=sys.stderr)
            print(f"  例: {_path} {_me} <出力先.xlsx>", file=sys.stderr)
            break
    raise SystemExit(1)

FONT = "Yu Gothic"          # 日本語の報告書向け。英語のみなら "Arial" に変える
NAVY = "1D3A5C"
GRAY = "5B6B7B"
EXAMPLE_FILL = "FFFBEA"     # 例示行の網掛け
LAST = 200                  # 数式が参照する最終行

# シート名は構成によって変わる。指摘事項一覧を参照する数式があるため、
# 名前をここで一元管理して、数式側にも同じものを使う。
# full の並びと番号（6b_ を含む）は、評価で使う台帳の構成に合わせてある。
LAYOUTS = {
    "none": {
        "summary": "1_総合評価", "runtime": "2_実機確認サマリ",
        "findings": "3_指摘事項一覧", "good": "4_評価できる実装",
        "unknown": "5_未確認事項", "roadmap": "6_対応ロードマップ",
    },
    "owasp": {
        "summary": "1_総合評価", "runtime": "2_実機確認サマリ",
        "findings": "3_指摘事項一覧", "good": "4_評価できる実装",
        "unknown": "5_未確認事項", "roadmap": "6_対応ロードマップ",
        "owasp": "7_枠組みへの当てはめ",
    },
    "full": {
        "summary": "1_総合評価", "runtime": "2_実機確認サマリ",
        "privacy": "3_個人情報セキュリティ", "owasp": "4_OWASP_Top10",
        "ipa": "5_IPA非機能要求グレード", "findings": "6_指摘事項一覧",
        "good": "6b_評価できる実装", "unknown": "7_残る未確認事項",
        "roadmap": "8_対応ロードマップ",
    },
}
# シートは LAYOUTS に書いた順に並ぶ（dict は挿入順を保つ）。
# 枠組みシートの位置が構成ごとに違う（owasp は最後、full は指摘一覧の前）ため、
# 全構成で共通の順序表を別に持つとここが必ずずれる。

# 指摘の 1 行には、判定・優先度・状態という別々の問いへの答えが入る（references/04-findings-register.md）。
# 見送りとクローズは優先度ではなく状態の値。クローズは「直して閉じた」と「もともと問題が無かった」を分ける。
VERDICTS = ["問題あり", "問題なし", "判断保留"]
PRIORITIES = ["P0", "P1", "P2", "P3", "P4", "—"]      # 問題なしの行は「—」
STATES = ["未対応", "対応中", "クローズ（解消）", "クローズ（該当なし）", "見送り"]
# 確認の方法は複数を併記することがある（「コード＋実機」）。一覧は候補として出すが、ほかの値も受け付ける。
METHODS = ["コード", "実機", "依頼者の確認", "取材", "コード＋実機"]
# 枠組み（個人情報・カード決済）への当てはめの判定。指摘の判定とは別の軸。
# コードから判定できない区分（人的・物理的など）は「範囲外」にして、未確認と混同させない。
OUT_OF_SCOPE = "範囲外（取材で聞く）"
FW_VERDICTS = ["適合", "不適合", "判断保留", OUT_OF_SCOPE]

PRIORITY_FILL = {
    "P0": "F5C6CB", "P1": "F8D7DA", "P2": "FDEBD0",
    "P3": "E3E9F3", "P4": "EDF1F6",
}

# 指摘事項一覧の列。集計の数式はここから列の位置を引くので、並べ替えてもずれない。
FINDING_COLS = [
    ("ID", 8), ("判定", 10), ("優先度", 8), ("状態", 14), ("分類", 16), ("確認の方法", 14),
    ("指摘事項", 80), ("該当箇所", 36), ("想定される影響", 80),
    ("AI実装(h)", 10), ("人手(h)", 10), ("是正案", 80),
]


def fcol(name):
    """指摘事項一覧で、その列の文字（A, B, …）を返す。"""
    for i, (h, _) in enumerate(FINDING_COLS, start=1):
        if h == name:
            return get_column_letter(i)
    raise KeyError(name)


thin = Side(style="thin", color="D6DCE4")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def title_block(ws, title, subtitle, width):
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT, size=9, color=GRAY)
    ws["A2"].alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.row_dimensions[2].height = 44


def header_row(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body_row(ws, row, values, example=False):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=FONT, size=10, italic=example)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BORDER
        if example:
            c.fill = PatternFill("solid", fgColor=EXAMPLE_FILL)


def note(ws, row, text, width):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=9, color=GRAY)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=width)


def section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = Font(name=FONT, size=11, bold=True, color=NAVY)


def add_list(ws, cells, values, strict=True):
    """セルに入力規則（一覧から選ぶ）を付ける。

    strict=False のときは一覧を候補として出すだけで、ほかの値も受け付ける。
    """
    dv = DataValidation(type="list", formula1='"{}"'.format(",".join(values)), allow_blank=True)
    dv.showDropDown = False           # False で一覧の矢印が出る（openpyxl の名前は逆の意味）
    dv.showErrorMessage = strict
    if strict:
        dv.errorTitle = "一覧にない値"
        dv.error = "次のいずれかを選ぶ: " + "／".join(values)
    ws.add_data_validation(dv)
    dv.add(cells)
    return dv


# --------------------------------------------------------------------------
def sheet_summary(wb, names, service, date, skill_version, standards):
    ws = wb.create_sheet(names["summary"])
    F = "'{}'".format(names["findings"])
    title_block(ws, f"{service} セキュリティ評価",
                f"評価日 {date}／個人情報を取り扱う Web アプリケーションとしての適合性評価。"
                "件数と工数は『{}』から自動集計される。".format(names["findings"]), 5)
    ws.column_dimensions["A"].width = 30
    for col in "BCDE":
        ws.column_dimensions[col].width = 22

    # 評価の前提。これが無いと、読み手は台帳の確度を判断できない（04 の「評価の前提」）。
    section(ws, 4, "■ 評価の前提")
    premises = [
        ("評価日", date, False),
        ("対象リビジョン", "（コミットハッシュ）", True),
        # 評価に使ったスキルの版。スクリプトは skill/ の外にある VERSION を読めないので、
        # --skill-version で渡されたときだけ入れる。渡されなければ空欄にして手で書く。
        ("評価に使ったスキルの版", skill_version or None, False),
        ("実機確認のモード", "（A：評価者が実機で確認／B：依頼者に実行してもらった／混在）", True),
        ("確認の範囲",
         "（コードで見たもの（全ハンドラの本数、走査した履歴の本数）／実機で見たもの／"
         "依頼者の確認で見たもの（返ってきたチェックリストと SQL）／見られなかったもの（U-x で示す））", True),
        ("照らした基準の版", standards, True),
        ("保持する個人情報", "（具体的に列挙。決済情報の保持有無を明記）", True),
        ("ログインする主体", "（役割ごとに、人数の桁と権限の範囲）", True),
        ("本評価の位置づけ",
         "本評価はコード監査と設定確認であり、脆弱性診断（動的な検査・ペネトレーションテスト）の"
         "代わりではない。", False),
    ]
    r = 5
    for k, v, placeholder in premises:
        ws.cell(row=r, column=1, value=k).font = Font(name=FONT, size=10, bold=True)
        c = ws.cell(row=r, column=2, value=v)
        c.font = Font(name=FONT, size=10, color=GRAY if placeholder else "000000")
        c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        if isinstance(v, str) and len(v) > 24:
            # 結合した B〜E 列の幅で、全角 24 文字ほどで折り返す
            ws.row_dimensions[r].height = 15 * (-(-len(v) // 24))
        r += 1

    # 対応が要る指摘。判定が 問題あり で、状態が クローズ でも 見送り でもないものを数える。
    # 「クローズ*」は クローズ（解消）と クローズ（該当なし）の両方に当たる。
    r += 1
    section(ws, r, "■ 対応が要る指摘（自動集計。判定が問題ありで、状態がクローズでも見送りでもないもの）")
    head = r + 1
    header_row(ws, head, ["優先度", "件数", "", "AI 実装(h)", "人手(h)"], [30, 10, 4, 12, 12])
    ws.freeze_panes = None

    def rng(col):
        c = fcol(col)
        return f"{F}!${c}$4:${c}${LAST}"

    open_crit = (f'{rng("判定")},"問題あり",{rng("状態")},"<>クローズ*",'
                 f'{rng("状態")},"<>見送り"')
    rows = [("P0（即日〜3日）", "P0"), ("P1（1〜2週間）", "P1"), ("P2（1か月）", "P2"),
            ("P3（3か月）", "P3"), ("P4（6か月）", "P4")]
    first = head + 1
    for i, (label, key) in enumerate(rows):
        rr = first + i
        body_row(ws, rr, [label, None, "件", None, None])
        crit = f'{open_crit},{rng("優先度")},"{key}"'
        ws.cell(row=rr, column=2).value = f"=COUNTIFS({crit})"
        ws.cell(row=rr, column=4).value = f'=SUMIFS({rng("AI実装(h)")},{crit})'
        ws.cell(row=rr, column=5).value = f'=SUMIFS({rng("人手(h)")},{crit})'
    rr = first + len(rows)
    body_row(ws, rr, ["合計", None, "件", None, None])
    for col in (2, 4, 5):
        L = get_column_letter(col)
        ws.cell(row=rr, column=col).value = f"=SUM({L}{first}:{L}{rr-1})"
        ws.cell(row=rr, column=col).font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=rr, column=1).font = Font(name=FONT, size=10, bold=True)

    # 台帳の内訳。サマリ・修正指示書・台帳で数え方を揃えるための但し書きに使う。
    r = rr + 2
    section(ws, r, "■ 台帳の内訳（数え方の但し書きに使う）")
    header_row(ws, r + 1, ["区分", "件数", ""], [30, 10, 4])
    ws.freeze_panes = None
    base = r + 2
    J, S, A = rng("判定"), rng("状態"), rng("ID")
    breakdown = [
        ("台帳の行数（ID のある行）", f"=COUNTA({A})"),
        ("問題あり", f'=COUNTIF({J},"問題あり")'),
        ("　うち 見送り", f'=COUNTIFS({J},"問題あり",{S},"見送り")'),
        ("　うち クローズ（解消）", f'=COUNTIFS({J},"問題あり",{S},"クローズ（解消）")'),
        ("判断保留", f'=COUNTIF({J},"判断保留")'),
        ("問題なし（クローズ（該当なし））", f'=COUNTIF({J},"問題なし")'),
        # 判定の書き忘れを見つけるための行。0 でなければ、どこかの行に判定が無い
        ("判定が空欄の行", f"=B{base}-B{base+1}-B{base+4}-B{base+5}"),
    ]
    for i, (label, formula) in enumerate(breakdown):
        body_row(ws, base + i, [label, formula, "件"])
    r = base + len(breakdown) + 1
    note(ws, r,
         "【数え方の但し書き】台帳 <n> 行 ＝ 問題あり <a> 行（うち 見送り <b>・クローズ（解消） <c>）"
         "＋ 判断保留 <d> 行 ＋ 問題なし <e> 行（クローズ（該当なし））。対応が要る指摘は <a−b−c> 件（P0〜P4）。"
         "別途、指摘ではない前提タスクとして T-1（<作業の名前>）がある。サマリと修正指示書にも同じものを書く。",
         5)
    r += 4
    note(ws, r,
         "【結論の書き方】土台として成立している点を先に置き、次に先に塞ぐべきものを 3 件まで。"
         "それぞれ「何が起きるか」を 2〜3 行で書く。5 件も 10 件も挙げるとどれが先か伝わらなくなる。",
         5)

    # 気づいたこと。評価の範囲外で目に付いたもの、事実として記録しておくだけのもの。
    # 指摘の表に入れず、判定も優先度も付けない。シートの末尾に置くので、行は下へ足せる。
    r += 4
    section(ws, r, "■ 気づいたこと（指摘にしないもの。判定も優先度も付けない）")
    header_row(ws, r + 1, ["区分", "内容"], [30, 22])
    ws.freeze_panes = None
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=5)
    body_row(ws, r + 2, ["記録", "保有件数は <n> 件、最古の登録日は <年月>（← 事実として残すだけのもの。"
                                 "評価の範囲外で目に付いたものもここに書く）"], example=True)
    ws.merge_cells(start_row=r + 2, start_column=2, end_row=r + 2, end_column=5)
    return ws


def sheet_runtime(wb, names):
    ws = wb.create_sheet(names["runtime"])
    title_block(ws, "実機確認サマリ",
                "コードからは判定できない実行環境の設定値について、参照系のみで確認した結果。"
                "データ変更・設定変更は行っていない。判定は 問題なし／問題あり／判断保留 の 3 値で、"
                "中間の値は作らない。指摘にしない気づいたことは、総合評価の「気づいたこと」に置く。", 5)
    header_row(ws, 4, ["区分", "確認項目", "結果（事実）", "判定", "関連ID"], [16, 30, 70, 12, 14])
    body_row(ws, 5, ["アクセス制御", "全テーブルの行レベル権限",
                     "全 <n> テーブルで有効。無効は 0 件（← 実行結果をそのまま書く。「問題ありませんでした」と要約しない）",
                     "問題なし", "—"], example=True)
    add_list(ws, f"D5:D{LAST}", VERDICTS)
    note(ws, 7,
         "【この欄の使い方】実行した SQL・コマンドの結果をそのまま残す。集約や整形をした場合は、"
         "元の行数と何をしたかを添える。例:「<n> 行返ったが読みづらいため、テーブル×ロールで集約"
         "（集約後 <m> 行、権限の総数は一致）」。", 5)
    return ws


def sheet_findings(wb, names):
    ws = wb.create_sheet(names["findings"])
    title_block(ws, "指摘事項一覧",
                "判定：問題あり／問題なし／判断保留。優先度：P0＝即日〜3日／P1＝1〜2週間／P2＝1か月／"
                "P3＝3か月／P4＝6か月（問題なしの行は「—」）。状態：未対応／対応中／クローズ（解消）／"
                "クローズ（該当なし）／見送り。ID は S-xx（コードを根拠にした指摘。外部の情報で、使っている版が"
                "当たると分かったものを含む）、N-xx（コードの外で判明し、裏が取れた指摘）。前提タスク（T-x）は"
                "この表に入れず、対応ロードマップに置く。一度振った ID は再利用しない。"
                "「指摘事項」には事実を、「想定される影響」には誰が何をできてしまうかを書く。", len(FINDING_COLS))
    ws.row_dimensions[2].height = 58
    header_row(ws, 3, [h for h, _ in FINDING_COLS], [w for _, w in FINDING_COLS])
    example = {
        "ID": "S-01", "判定": "問題あり", "優先度": "P0", "状態": "未対応", "分類": "認可",
        "確認の方法": "コード",
        "指摘事項": "注文の詳細を返すハンドラが、ログインの有無は確かめるが、注文の持ち主が要求した本人かを"
                    "確かめていない。パスの注文 ID を書き換えると、他人の注文の氏名・配送先住所が返る。",
        "該当箇所": "<ファイル>:<行>",
        "想定される影響": "会員登録をした誰でも、注文 ID を順に変えて全利用者の氏名と住所を取得できる。"
                          "ID は連番なので総当たりの手間も無い。",
        "AI実装(h)": 1, "人手(h)": 0.5,
        "是正案": "注文を取り出すときに、要求した利用者を条件に含める（持ち主で絞る）。同じ形のハンドラを"
                  "一覧にして、同じ直し方を当てる。",
    }
    body_row(ws, 4, [example[h] for h, _ in FINDING_COLS], example=True)
    ws[f"{fcol('優先度')}4"].fill = PatternFill("solid", fgColor=PRIORITY_FILL["P0"])
    add_list(ws, f"{fcol('判定')}4:{fcol('判定')}{LAST}", VERDICTS)
    add_list(ws, f"{fcol('優先度')}4:{fcol('優先度')}{LAST}", PRIORITIES)
    add_list(ws, f"{fcol('状態')}4:{fcol('状態')}{LAST}", STATES)
    add_list(ws, f"{fcol('確認の方法')}4:{fcol('確認の方法')}{LAST}", METHODS, strict=False)
    # 合計行はこのシートに置かない。同一シート内に置くと、集計範囲を広く取ったときに
    # 合計セル自身を巻き込んで循環参照になる。範囲を狭く取れば今度は行を足したときに
    # 数式へ手を入れる必要が出る。集計は総合評価シートに一本化してある。
    return ws


def sheet_good(wb, names):
    ws = wb.create_sheet(names["good"])
    title_block(ws, "評価できる実装（維持すべき点）",
                "今後の改修で壊してはいけない実装。これを書かないと、修正の過程で壊れる。"
                "実機確認で裏が取れたものには【実機】と印を付ける。", 2)
    header_row(ws, 4, ["区分", "内容"], [20, 110])
    body_row(ws, 5, ["認可",
                     "全 API ルートにガードが存在する。無認証は問い合わせ受付・計測・ログアウト等、"
                     "意図された公開口のみ【実機】"], example=True)
    return ws


def sheet_unknown(wb, names):
    ws = wb.create_sheet(names["unknown"])
    title_block(ws, "残る未確認事項",
                "コード監査と実機確認を経て、なお確定していない項目。"
                "重要度「高」は、何かの作業を止めているもの。"
                "「調べれば分かるもの」と「担当者しか知らないもの」を区別して書く。", 5)
    header_row(ws, 4, ["No", "確認したい情報", "重要度", "取得方法", "影響する項目"], [8, 44, 10, 70, 16])
    body_row(ws, 5, ["U-1", "管理者アカウントの一覧（件数と、それぞれの持ち主）", "高",
                     "担当者への確認が唯一の手段。認証基盤の管理画面の閲覧権限が評価者に無く、"
                     "API からも一覧を読み出せないため。",
                     "S-02"], example=True)
    note(ws, 7,
         "【作業を止めているものは表の外にも書く】例:「U-1 は S-02 の作業を止める。"
         "担当者への確認以外に取得手段がないため、着手と同時に依頼を出すこと。」", 5)
    return ws


def sheet_roadmap(wb, names):
    ws = wb.create_sheet(names["roadmap"])
    title_block(ws, "対応ロードマップ",
                "依存関係で並べる。深刻度順ではない。前提になっているもの（プラン変更、調査タスク、"
                "スキーマ復元）を先に置く。指摘ではない前提タスクは T-x の ID でここに置く"
                "（指摘事項一覧には入れない）。工数に待ち時間（確認待ち・観測期間・承認）は含めない。", 7)
    header_row(ws, 4, ["フェーズ", "時期", "対象ID", "実施内容", "AI実装", "人手", "補足"],
               [13, 13, 15, 44, 11, 11, 88])
    body_row(ws, 5, ["フェーズ0", "即日", "T-1",
                     "（前提タスク）修正の前に、本番のデータを戻せる状態にする（プランの変更を含む）", 0, 1,
                     "指摘ではないが、以降の作業の前提になる。指摘の件数には数えない"], example=True)
    body_row(ws, 6, ["フェーズ0", "即日〜3日", "S-01", "注文の詳細を返すハンドラに持ち主の確認を足す", 1, 0.5,
                     "同じ形のハンドラを一覧にして、まとめて直す"], example=True)
    return ws


# 個人情報の安全管理措置のチェックリスト。
# 区分は個人情報保護委員会ガイドライン（通則編）の「10（別添）講ずべき安全管理措置の内容」の 7 区分に揃える。
# 対外説明で「安全管理措置を講じているか」を問われたとき、ガイドラインの区分に沿って答えられるようにするため
# （references/06-frameworks.md の「安全管理措置の 7 区分」）。観点は実装上の定石で、案件に合わせて足し引きする。
PRIVACY_AREAS = [
    "基本方針の策定", "取扱いに係る規律の整備", "組織的安全管理措置", "人的安全管理措置",
    "物理的安全管理措置", "技術的安全管理措置", "外的環境の把握",
]
# コードから判定できない区分。既定の判定を「範囲外（取材で聞く）」にして、未確認と混同させない。
PRIVACY_OUT_OF_SCOPE = ("人的安全管理措置", "物理的安全管理措置")
PRIVACY_ITEMS = [
    ("基本方針の策定", "基本方針", "個人情報の取扱いに関する基本方針（プライバシーポリシー）を定め、公表しているか"),
    ("基本方針の策定", "文書と実装", "公表している取得項目・利用目的・第三者提供と、実装で取得・送信している項目が一致しているか"),
    ("取扱いに係る規律の整備", "取扱規程", "取得・利用・保存・提供・削除の手順が定められているか"),
    ("取扱いに係る規律の整備", "保存と削除", "個人情報の保存期間・削除方針が定義されているか"),
    ("取扱いに係る規律の整備", "第三者提供", "個人情報の第三者提供が制御されているか"),
    ("取扱いに係る規律の整備", "ログ", "ログの保存期間が定義されているか"),
    ("組織的安全管理措置", "体制", "個人情報の取扱いの責任者と、報告連絡の体制が決まっているか"),
    ("組織的安全管理措置", "取扱状況の把握", "個人情報へのアクセス記録を保存しているか"),
    ("組織的安全管理措置", "取扱状況の把握", "異常検知・アラートの仕組みがあるか"),
    ("組織的安全管理措置", "漏えい時の対応", "インシデント対応手順があるか（漏えい等の報告の期限を含む）"),
    ("組織的安全管理措置", "委託先の監督", "委託先の安全管理措置を確認しているか"),
    ("組織的安全管理措置", "委託先の監督", "委託先に渡す権限が必要最小限に絞られているか"),
    ("人的安全管理措置", "教育・監督", "従業者に個人情報の取扱いを教育し、監督しているか"),
    ("物理的安全管理措置", "区域と機器", "個人情報を取り扱う区域と、機器・電子媒体を管理しているか（マネージド基盤なら基盤側の責任範囲）"),
    ("技術的安全管理措置", "認証", "認証基盤に実績のある仕組みを使っているか"),
    ("技術的安全管理措置", "認証", "管理者権限に多要素認証が設定されているか"),
    ("技術的安全管理措置", "認証", "パスワードポリシーが定義されているか"),
    ("技術的安全管理措置", "認証", "パスワードの配布方法が安全か"),
    ("技術的安全管理措置", "認証", "パスワード再設定の導線が安全か"),
    ("技術的安全管理措置", "認証", "アカウント発行が制御されているか"),
    ("技術的安全管理措置", "認証", "クライアント側のみの認証判定に依存していないか"),
    ("技術的安全管理措置", "認証", "トークンの署名を検証しているか（復号だけで済ませていないか）"),
    ("技術的安全管理措置", "認証", "外部の認証基盤を使う場合、戻り先 URL が完全一致で設定されているか"),
    ("技術的安全管理措置", "認可", "全ての API エンドポイントに認可判定があるか"),
    ("技術的安全管理措置", "認可", "ロール情報をクライアントが改竄できない場所に保持しているか"),
    ("技術的安全管理措置", "認可", "他人のデータを ID 指定で操作できないか（IDOR）"),
    ("技術的安全管理措置", "認可", "認可の付け忘れを機械的に検出できるか"),
    ("技術的安全管理措置", "認可", "認可の枠組み自体が破られた場合の備えがあるか（多層になっているか）"),
    ("技術的安全管理措置", "認可", "リアルタイム通信（WebSocket・購読・チャネル）に、購読ごとの認可があるか（07 の 11 節）"),
    ("技術的安全管理措置", "データ保護", "行レベルの権限制御が有効か"),
    ("技術的安全管理措置", "データ保護", "公開鍵から業務データが読めない設計か"),
    ("技術的安全管理措置", "データ保護", "権限制御の対象外となる権限が残っていないか"),
    ("技術的安全管理措置", "データ保護", "権限制御を迂回する鍵がクライアントへ流出しない仕組みがあるか"),
    ("技術的安全管理措置", "データ保護", "権限制御を迂回する経路が無いか（定義者権限の関数・ビュー）"),
    ("技術的安全管理措置", "データ保護", "BaaS の設定（Supabase の RLS と GRANT、Firebase のルール）で、匿名・一般利用者の鍵から業務データへ届かないか（02 の E 節、03 の 1 節）"),
    ("技術的安全管理措置", "データ保護", "画面操作の記録（セッションリプレイ）で、入力値と画面上の個人情報がマスクされているか（08 の 1-2）"),
    ("技術的安全管理措置", "データ保護", "通信が暗号化されているか"),
    ("技術的安全管理措置", "データ保護", "保管データが暗号化されているか"),
    ("技術的安全管理措置", "データ保護", "DB への到達経路が制限されているか"),
    ("技術的安全管理措置", "入力検証", "SQL インジェクション対策"),
    ("技術的安全管理措置", "入力検証", "サーバー側での入力検証を行っているか"),
    ("技術的安全管理措置", "入力検証", "XSS 対策"),
    ("技術的安全管理措置", "入力検証", "CSV 出力の数式インジェクション対策"),
    ("技術的安全管理措置", "入力検証", "SSRF 対策"),
    ("技術的安全管理措置", "濫用対策", "公開エンドポイントに BOT 対策があるか"),
    ("技術的安全管理措置", "濫用対策", "レート制限があるか"),
    ("技術的安全管理措置", "濫用対策", "メール送信が第三者に悪用されないか"),
    ("技術的安全管理措置", "濫用対策", "業務上の課金経路が保護されているか"),
    ("技術的安全管理措置", "濫用対策", "従量課金の経路に上限と監視があるか"),
    ("技術的安全管理措置", "濫用対策", "SMS の送信経路（確認コード・通知）に、上限・送信先の国の制限・監視があるか（02 の F-4）"),
    ("技術的安全管理措置", "秘密情報", "ソースコードに秘密情報が含まれていないか"),
    ("技術的安全管理措置", "秘密情報", "秘密情報の保管場所が適切か"),
    ("技術的安全管理措置", "秘密情報", "鍵の用途が分離されているか"),
    ("技術的安全管理措置", "秘密情報", "使用していない鍵が無効化されているか"),
    ("技術的安全管理措置", "構成管理", "セキュリティヘッダが設定されているか"),
    ("技術的安全管理措置", "構成管理", "本番に開発用のバイパスが残っていないか"),
    ("技術的安全管理措置", "構成管理", "検証環境が本番から分離されているか"),
    ("技術的安全管理措置", "構成管理", "Webhook の真正性を検証しているか"),
    ("技術的安全管理措置", "構成管理", "例外が起きたときに権限が開く側へ倒れないか"),
    ("技術的安全管理措置", "構成管理", "DB スキーマがバージョン管理されているか"),
    ("技術的安全管理措置", "構成管理", "静的解析が CI で動いているか"),
    ("技術的安全管理措置", "構成管理", "CI の定義が、外部からのプルリクエストで秘密情報や書き込み権限に届かないか（10 の 3-3）"),
    ("技術的安全管理措置", "構成管理", "AI エージェントの設定（AGENTS.md・.mcp.json など）が、秘密情報の読み出しや任意のコマンド実行を許していないか（10 の 3-5）"),
    ("技術的安全管理措置", "脆弱性管理", "依存パッケージの脆弱性を継続監視しているか"),
    ("技術的安全管理措置", "脆弱性管理", "依存の取得元と名前を確認しているか"),
    ("技術的安全管理措置", "脆弱性管理", "基盤のバージョンが最新に保たれているか"),
    ("技術的安全管理措置", "脆弱性管理", "定期的な脆弱性診断を実施しているか"),
    ("技術的安全管理措置", "事業継続", "バックアップが取得されているか"),
    ("技術的安全管理措置", "事業継続", "復旧目標（RTO / RPO）が定義されているか"),
    ("外的環境の把握", "越境移転", "外国にある事業者（クラウド・SaaS）に個人データを預けている場合、その国の制度を把握しているか"),
    ("外的環境の把握", "越境移転", "越境移転について本人に説明できるか"),
]

# IPA 非機能要求グレードの 6 大項目と、評価でよく使う中項目。
IPA_METRICS = [
    ("A. 可用性", "継続性：運用スケジュール（サービス時間）"),
    ("A. 可用性", "継続性：目標復旧水準（RTO / RPO）"),
    ("A. 可用性", "継続性：稼働率"),
    ("A. 可用性", "耐障害性：サーバ／データの冗長化"),
    ("A. 可用性", "回復性：バックアップ方式・保存期間"),
    ("A. 可用性", "災害対策：システム復旧（リージョン障害）"),
    ("B. 性能・拡張性", "業務量：通常時／ピーク時のリクエスト数"),
    ("B. 性能・拡張性", "性能目標値：レスポンスタイム"),
    ("B. 性能・拡張性", "リソース拡張性：CPU / メモリ / ディスク"),
    ("B. 性能・拡張性", "性能試験：試験実施の有無"),
    ("C. 運用・保守性", "通常運用：運用時間・バックアップ運用"),
    ("C. 運用・保守性", "通常運用：監視（稼働・性能・異常）"),
    ("C. 運用・保守性", "保守運用：パッチ適用（脆弱性対応）"),
    ("C. 運用・保守性", "運用環境：ログ取得・保存期間"),
    ("C. 運用・保守性", "運用環境：検証環境の分離"),
    ("C. 運用・保守性", "運用体制：障害時の連絡体制"),
    ("D. 移行性", "移行方式：移行スケジュール・体制"),
    ("D. 移行性", "移行データ：移行対象・移行ツール"),
    ("E. セキュリティ", "前提条件：情報資産の重要度"),
    ("E. セキュリティ", "認証：本人認証（認証方式）"),
    ("E. セキュリティ", "利用制限：管理方法（アクセス制御）"),
    ("E. セキュリティ", "データの秘匿：暗号化（通信／保管）"),
    ("E. セキュリティ", "不正監視：監視対象・追跡（監査ログ）"),
    ("E. セキュリティ", "Web 対策：不正アクセス対策"),
    ("E. セキュリティ", "ネットワーク対策：経路の制限"),
    ("E. セキュリティ", "マルウェア対策・脆弱性管理"),
    ("E. セキュリティ", "セキュリティリスク分析・診断"),
    ("F. システム環境・エコロジー", "システム制約・前提条件"),
    ("F. システム環境・エコロジー", "環境マネジメント・耐震／免震"),
]


def sheet_privacy(wb, names):
    ws = wb.create_sheet(names["privacy"])
    title_block(ws, "個人情報取扱いのチェックリスト（安全管理措置）",
                "区分は個人情報保護委員会ガイドライン（通則編）「10（別添）講ずべき安全管理措置の内容」の 7 区分。"
                "判定は 適合／不適合／判断保留 の 3 値（指摘の判定とは別の軸）。コードから判定できない区分"
                "（人的・物理的）は「範囲外（取材で聞く）」にし、未確認と混同させない。根拠には実行結果を書き、"
                "実機で確定したものには【実機確定】と印を付ける。関連 ID が無い行は「—」。"
                "技術的安全管理措置だけを見て「安全管理措置を講じている」とは書かない。", 6)
    ws.row_dimensions[2].height = 58
    header_row(ws, 4, ["区分", "観点", "確認項目", "判定", "確認結果（根拠）", "関連ID"],
               [22, 14, 50, 14, 70, 12])
    body_row(ws, 5, ["技術的安全管理措置", "認証", "管理者権限に多要素認証が設定されているか", "不適合",
                     "【実機確定】管理者 <n> アカウントのうち、多要素認証の登録は 0 件（← 実行結果を"
                     "そのまま書く。「設定されていませんでした」と要約しない）",
                     "S-02"], example=True)
    for i, (area, view, item) in enumerate(PRIVACY_ITEMS, start=6):
        body_row(ws, i, [area, view, item, OUT_OF_SCOPE if area in PRIVACY_OUT_OF_SCOPE else "", "", ""])
    add_list(ws, f"D5:D{LAST}", FW_VERDICTS)
    note(ws, len(PRIVACY_ITEMS) + 7,
         "【未確認と判断保留を区別する】「調べれば分かる」ものは未確認事項（U-x）に回す。"
         "「事業側が決める」ものが判断保留。前者を判断保留に混ぜると、誰も動かないまま残る。"
         "【範囲外は取材で埋める】人的・物理的の区分はコードから見えない。取材で聞けたら、聞いた内容と"
         "「取材による」ことを根拠に書く（取材だけを根拠に適合としない）。", 6)
    return ws


def sheet_ipa(wb, names):
    ws = wb.create_sheet(names["ipa"])
    title_block(ws, "IPA 非機能要求グレードによる水準評価",
                "6 大項目について、現状の水準と事業として求められる水準を並べ、ギャップを示す。"
                "可用性や性能が低い水準にあること自体は、事業フェーズによっては合理的な選択。"
                "この表の価値は、セキュリティだけが低い水準に留まってはいけない理由を示せる点にある。", 6)

    ws["A4"] = "■ モデルシステムとの対比"
    ws["A4"].font = Font(name=FONT, size=11, bold=True, color=NAVY)
    header_row(ws, 5, ["区分", "定義", "本システムとの関係", "", "内容", ""], [22, 32, 22, 4, 60, 4])
    ws.merge_cells(start_row=5, start_column=3, end_row=5, end_column=4)
    ws.merge_cells(start_row=5, start_column=5, end_row=5, end_column=6)
    ws.freeze_panes = None
    for i, (a, b) in enumerate([
        ("モデルシステム①", "社会的影響が殆ど無いシステム"),
        ("モデルシステム②", "社会的影響が限定されるシステム"),
        ("モデルシステム③", "社会的影響が極めて大きいシステム"),
    ], start=6):
        body_row(ws, i, [a, b, "（現状水準／求められる水準／該当しない）", "", "", ""])
        ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=4)
        ws.merge_cells(start_row=i, start_column=5, end_row=i, end_column=6)

    ws["A10"] = "■ 大項目別の評価"
    ws["A10"].font = Font(name=FONT, size=11, bold=True, color=NAVY)
    header_row(ws, 11, ["大項目", "中項目（メトリクス）", "現状Lv", "推奨Lv",
                        "現状の根拠", "ギャップを埋める施策"], [22, 42, 9, 9, 66, 66])
    ws.freeze_panes = None
    body_row(ws, 12, ["A. 可用性", "回復性：バックアップ方式・保存期間", 0, 3,
                      "【実機確定】日次バックアップも時点復旧も存在しない（← 実行結果をそのまま書く）",
                      "有償プランへ移行し日次バックアップを有効化。復元テストまで行う"],
             example=True)
    for i, (dai, chu) in enumerate(IPA_METRICS, start=13):
        body_row(ws, i, [dai, chu, "", "", "", ""])
    note(ws, len(IPA_METRICS) + 14,
         "【全項目を上げましょう、と書かない】上げるべき 1 項目を特定して、"
         "そのギャップの中身を 4 つ以内に絞る。個人情報を第三者に開示する業務を担っているなら、"
         "他の項目が低水準でも、セキュリティだけは一段上が要る。この対比が投資判断の材料になる。", 6)
    return ws


# OWASP Top 10 の版。既定は 2025。依頼者が既に 2021 版で文書を作っている場合や、
# 取引先の審査項目が 2021 版で書かれている場合のために、2021 版も選べるようにしてある。
OWASP = {
    "2025": [
        ("A01", "Broken Access Control"),
        ("A02", "Security Misconfiguration"),
        ("A03", "Software Supply Chain Failures"),
        ("A04", "Cryptographic Failures"),
        ("A05", "Injection"),
        ("A06", "Insecure Design"),
        ("A07", "Authentication Failures"),
        ("A08", "Software or Data Integrity Failures"),
        ("A09", "Security Logging & Alerting Failures"),
        ("A10", "Mishandling of Exceptional Conditions"),
    ],
    "2021": [
        ("A01", "Broken Access Control"),
        ("A02", "Cryptographic Failures"),
        ("A03", "Injection"),
        ("A04", "Insecure Design"),
        ("A05", "Security Misconfiguration"),
        ("A06", "Vulnerable and Outdated Components"),
        ("A07", "Identification and Authentication Failures"),
        ("A08", "Software and Data Integrity Failures"),
        ("A09", "Security Logging and Monitoring Failures"),
        ("A10", "Server-Side Request Forgery"),
    ],
}


# OWASP API Security Top 10 (2023)。API が主体の構成のときだけ使う。
# 一般の Top 10 と重なる部分が多いので、両方を並べない（--api は OWASP のシートを置き換える）。
API_TOP10 = [
    ("API1", "Broken Object Level Authorization", "他人の ID を指定して到達できないか（監査 A-3）"),
    ("API2", "Broken Authentication", "認証の強度、トークンの検証（監査 B 群）"),
    ("API3", "Broken Object Property Level Authorization",
     "更新でオブジェクトをそのまま渡していないか／応答に不要な項目が無いか"),
    ("API4", "Unrestricted Resource Consumption", "レート制限、件数の上限、ページングの強制（監査 F-2）"),
    ("API5", "Broken Function Level Authorization", "管理用の操作を一般利用者が呼べないか（監査 A-1）"),
    ("API6", "Unrestricted Access to Sensitive Business Flows",
     "自動化されると困る流れに対策があるか（監査 F-1）"),
    ("API7", "Server Side Request Forgery", "外部への通信先を外部入力から組み立てていないか"),
    ("API8", "Security Misconfiguration", "設定の既定値、開発用の残骸（監査 I ／ 実機確認）"),
    ("API9", "Improper Inventory Management", "使われていない旧版の API が生きていないか（監査 J）"),
    ("API10", "Unsafe Consumption of APIs", "呼び出している外部 API の応答を検証しているか"),
]


def sheet_api(wb, names):
    ws = wb.create_sheet(names["api"])
    title_block(ws, "OWASP API Security Top 10 (2023) への当てはめ",
                "API が主体の構成のときに使う。一般の Top 10 と重なるため、両方は並べない。"
                "一般の Top 10 では拾いにくい API3（プロパティ単位の認可）と "
                "API9（旧版 API の放置）のために当てはめる価値がある。", 5)
    header_row(ws, 4, ["No", "カテゴリ", "判定", "主に見るもの", "根拠（要約）"],
               [9, 46, 14, 60, 90])
    for i, (no, name, watch) in enumerate(API_TOP10, start=5):
        body_row(ws, i, [no, name, "", watch, ""])
    note(ws, len(API_TOP10) + 6,
         "【API3 と API9 を書き落とさない】API3 は応答そのものを見ないと分からない"
         "（画面で使っていない項目が含まれていないか）。API9 は旧版が残っていると、"
         "新版で塞いだ穴がそちらに残る。ルート一覧に版ごとの重複が無いかを見る。", 5)
    return ws


def sheet_owasp(wb, names, edition="2025"):
    ws = wb.create_sheet(names["owasp"])
    title_block(ws, f"OWASP Top 10 ({edition}) への当てはめ",
                f"OWASP Top 10 の {edition} 版で判定する。個別の指摘を洗い出した後に行う。"
                "対外説明・網羅性の確認が要る場合のみ。判定だけでなく根拠を 1 行添える。"
                "「適合」だけだと、見たうえでの適合なのか、見ていないだけなのかが分からない。", 4)
    header_row(ws, 4, ["No", "カテゴリ", "判定", "根拠（要約）"], [8, 46, 16, 100])
    rows = OWASP[edition]
    for i, (no, name) in enumerate(rows, start=5):
        body_row(ws, i, [no, name, "", ""])
    extra = ("【どの版で書いたかを必ず明記する】版を書かない適合表は、読み手が何と比べているか"
             "分からない。2025 版では SSRF が A01 に統合され、A03 がサプライチェーン全体に、"
             "A10 が例外処理（fail-open を含む）になっている。"
             if edition == "2025" else
             "【どの版で書いたかを必ず明記する】これは 2021 版。2025 版では順序と範囲が変わっている"
             "（SSRF は A01 に統合、A03 はサプライチェーン全体、A10 は例外処理）。"
             "新しく作る報告書では 2025 版を使う。")
    note(ws, len(rows) + 6,
         extra + "【静的検証の限界を明記する】動的な検査（実際に攻撃を試す）をしていないなら、"
         "そう書く。判定できる範囲がどこまでかを示すことが、報告書の信頼性になる。", 4)
    return ws


# カード決済を扱うとき（決済代行の画面へ遷移する・埋め込む構成を含む）だけ作るシート。
# 常設しない。カード決済の無い案件に並べると、見ていない項目が「確認済み」に見える。
# 中身は references/06-frameworks.md の「カード決済を扱う場合」に揃える。
CARD_ITEMS = [
    ("PCI DSS v4.0.1", "6.4.3",
     "決済ページで動くスクリプトを一覧にし、1 本ずつ必要性を承認し、完全性を確かめている",
     "02 の L 節、10 の 3-4、09 の 1 節・10 節", "", ""),
    ("PCI DSS v4.0.1", "11.6.1",
     "決済ページとその HTTP ヘッダの改ざんを検知し、担当者に知らせる仕組みがある",
     "02 の L 節、09 の 1 節・10 節", "", ""),
    ("PCI DSS SAQ A（2025-01 改訂）", "適格要件",
     "iframe 型の決済で、サイトがスクリプトによる攻撃の影響を受けないことを確かめている"
     "（自社で 6.4.3 / 11.6.1 相当を講じるか、決済代行の保証を得る）",
     "決済の iframe を置いているページの CSP と第三者スクリプト", "", ""),
    ("クレジットカード・セキュリティガイドライン 6.1 版", "脆弱性対策 1（管理画面）",
     "管理画面の IP 制限と二段階認証、ログイン失敗 10 回以下でのロック", "02 の B-3", "", ""),
    ("クレジットカード・セキュリティガイドライン 6.1 版", "脆弱性対策 2（公開ディレクトリ）",
     "公開ディレクトリに重要なファイルを置かない", "03 の 10 節（外から見えてはいけないもの）", "", ""),
    ("クレジットカード・セキュリティガイドライン 6.1 版", "脆弱性対策 3（脆弱性診断）",
     "脆弱性診断またはペネトレーションテストの定期実施（自社開発ならソースコードレビューも）",
     "本評価は診断の代わりにならない", "",
     "本評価はコード監査と設定確認であり、ガイドラインが求める脆弱性診断の実施を示すものではない"),
    ("クレジットカード・セキュリティガイドライン 6.1 版", "脆弱性対策 4（ウイルス対策）",
     "ウイルス対策", "コードからは見えない", OUT_OF_SCOPE, ""),
    ("クレジットカード・セキュリティガイドライン 6.1 版", "脆弱性対策 5（クレジットマスター）",
     "クレジットマスター対策", "02 の F-5", "", ""),
    ("クレジットカード・セキュリティガイドライン 6.1 版", "EMV 3-D セキュア",
     "EMV 3-D セキュアの導入（ガイドラインが求めた期限は 2025 年 3 月末。過ぎている）",
     "決済代行の管理画面か契約（依頼者に確かめてもらう）", "", ""),
]


def sheet_card(wb, names):
    ws = wb.create_sheet(names["card"])
    title_block(ws, "カード決済を扱う場合の確認",
                "カード決済を扱うとき（カード番号を自社で持たず、決済代行の画面へ遷移する・埋め込む構成を含む）だけ作る。"
                "決済画面を読み込む自社ページのスクリプトが改ざんされれば、入力されたカード番号は抜かれる。"
                "判定は 適合／不適合／判断保留（指摘の判定とは別の軸）。コードから見えないものは「範囲外（取材で聞く）」。", 7)
    ws.row_dimensions[2].height = 58
    header_row(ws, 4, ["基準", "項目", "求めていること", "このスキルで見るところ", "判定", "確認結果（根拠）", "関連ID"],
               [26, 24, 56, 34, 14, 60, 10])
    for i, row in enumerate(CARD_ITEMS, start=5):
        body_row(ws, i, list(row) + [""])
    add_list(ws, f"E5:E{LAST}", FW_VERDICTS)
    note(ws, len(CARD_ITEMS) + 6,
         "【期限を過ぎた未実施は事実として書く】EMV 3-D セキュアと PCI DSS の 6.4.3 / 11.6.1 は期限を過ぎている。"
         "未実施なら「期限を過ぎて実施されていない」と指摘事項に書き、攻撃経路で優先度を引き直す。"
         "経路が描けなくても P2 より下に置かない（references/04 の「期限を過ぎた法令・基準の要求」）。"
         "【決済ページに広告・計測タグが同居していれば、それだけで指摘になる】第三者スクリプトは、"
         "カード番号の入力欄と同じ文書で動く。", 7)
    return ws


# --------------------------------------------------------------------------
def next_number(names):
    """任意のシートを末尾に足すときの番号。枚数ではなく、使っている番号の最大 + 1。

    full は「6b_」のように番号を枝分かれさせているので、枚数から数えると番号が飛ぶ。
    """
    nums = [int(v.split("_", 1)[0]) for v in names.values() if v.split("_", 1)[0].isdigit()]
    return max(nums) + 1


def main():
    ap = argparse.ArgumentParser(description="指摘台帳（xlsx）の雛形を生成する")
    ap.add_argument("output", help="出力先の .xlsx パス")
    ap.add_argument("--service", default="（サービス名）", help="対象サービス名")
    ap.add_argument("--date", default=_dt.date.today().isoformat(), help="評価日 (YYYY-MM-DD)")
    ap.add_argument("--skill-version", default="",
                    help="評価に使ったスキルの版（配布物の VERSION の値）。省くと総合評価の欄は空欄になる")
    ap.add_argument("--frameworks", choices=("none", "owasp", "full"), default="owasp",
                    help="枠組みへの当てはめシートの構成（既定: owasp）")
    ap.add_argument("--owasp", choices=("2025", "2021"), default="2025",
                    help="OWASP Top 10 の版（既定: 2025）")
    ap.add_argument("--api", action="store_true",
                    help="OWASP のシートを API Security Top 10 のシートに置き換える（API が主体の構成のとき）")
    ap.add_argument("--card", action="store_true",
                    help="カード決済を扱う場合のシートを足す（PCI DSS 6.4.3 / 11.6.1、SAQ A、"
                         "クレジットカード・セキュリティガイドライン、EMV 3-D セキュア）")
    ap.add_argument("--force", action="store_true",
                    help="出力先が既にあれば上書きする（既定では止まる）")
    args = ap.parse_args()

    # 書きかけの台帳を雛形で黙って潰さない。拡張子が違えば、表計算ソフトが開けない
    # ファイルができる（.xls や拡張子なしで保存しても中身は xlsx のまま）。
    if not args.output.lower().endswith(".xlsx"):
        ap.error(f"出力先の拡張子は .xlsx にする: {args.output}")
    if os.path.exists(args.output) and not args.force:
        print(f"既にある: {args.output}", file=sys.stderr)
        print("  書きかけの台帳を雛形で上書きしないよう止めた。上書きするなら --force を付ける。",
              file=sys.stderr)
        raise SystemExit(1)

    names = dict(LAYOUTS[args.frameworks])
    if args.api:
        # 一般の Top 10 と重なるため、両方は並べない（references/06-frameworks.md）。
        # OWASP のシートがある構成では、その位置と番号のまま API のシートに置き換える。
        # dict は挿入順で並ぶため、作り直して位置を保つ。
        if "owasp" in names:
            rebuilt = {}
            for k, v in names.items():
                if k == "owasp":
                    rebuilt["api"] = v.split("_", 1)[0] + "_API_Top10"
                else:
                    rebuilt[k] = v
            names = rebuilt
        else:
            names["api"] = f"{next_number(names)}_API_Top10"
    if args.card:
        names["card"] = f"{next_number(names)}_カード決済"

    standards = "（references/06-frameworks.md の表から、照らした基準と版を書く）"
    if "owasp" in names:
        standards = f"（OWASP Top 10 {args.owasp} ほか。references/06-frameworks.md の表から書く）"
    elif "api" in names:
        standards = "（OWASP API Security Top 10 2023 ほか。references/06-frameworks.md の表から書く）"

    builders = {
        "summary": lambda: sheet_summary(wb, names, args.service, args.date,
                                         args.skill_version, standards),
        "runtime": lambda: sheet_runtime(wb, names),
        "privacy": lambda: sheet_privacy(wb, names),
        "owasp": lambda: sheet_owasp(wb, names, args.owasp),
        "api": lambda: sheet_api(wb, names),
        "ipa": lambda: sheet_ipa(wb, names),
        "findings": lambda: sheet_findings(wb, names),
        "good": lambda: sheet_good(wb, names),
        "unknown": lambda: sheet_unknown(wb, names),
        "roadmap": lambda: sheet_roadmap(wb, names),
        "card": lambda: sheet_card(wb, names),
    }

    wb = Workbook()
    wb.remove(wb.active)
    for key in names:
        builders[key]()
    wb.save(args.output)

    suffix = f" / OWASP {args.owasp} 版" if "owasp" in names else ""
    if "api" in names:
        suffix += " / OWASP API Security Top 10 (2023)"
    print(f"生成: {args.output}（--frameworks {args.frameworks}{suffix}）")
    print(f"  シート {len(wb.sheetnames)} 枚: {', '.join(wb.sheetnames)}")
    print()
    print("  ・薄い黄色の斜体行は書き方の例。実際の指摘を書くときに上書きするか削除する")
    print("  ・判定・優先度・状態は一覧から選ぶ（入力規則）。見送りとクローズは優先度ではなく状態に書く")
    print("  ・件数と工数の合計は数式。行を足せば自動で追従する（参照範囲は 200 行まで）")
    print("  ・表計算ソフトで開くまで合計欄は空に見える。openpyxl は計算結果を持たないため")
    if not args.skill_version:
        print("  ・評価に使ったスキルの版は空欄。--skill-version で渡すか、総合評価に手で書く")


if __name__ == "__main__":
    main()
